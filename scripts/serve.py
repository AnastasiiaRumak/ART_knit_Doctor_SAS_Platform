"""
serve.py — Локальный веб-сервер ArtKnit.

Архитектура доступа:

  ПУБЛИЧНЫЕ:
    GET  /                       → редирект на /upload
    GET  /login                  → страница входа (4 роли)
    GET  /upload                 → форма загрузки протоколов
    POST /api/login              → авторизация (логин+пароль)
    GET  /api/me                 → проверка сессии
    GET  /api/logout             → выход (303 redirect на /upload)
    POST /api/upload             → обработка протокола
    GET  /patient/{card_id}      → magic-link пациента (без пароля)

  ЗАЩИЩЁННЫЕ:
    GET  /patient/dashboard      → patient, admin
    GET  /doctor/dashboard       → doctor, admin
    GET  /doctor/upload          → doctor, admin
    GET  /doctor/{card_id}       → doctor, admin
    GET  /manager                → manager, admin
    GET  /admin                  → admin

  ⚠️ Конкретные пути объявлены ДО параметрических.

Запуск:
    python3 scripts/serve.py
    uvicorn scripts.serve:app --reload --port 8000
"""

from __future__ import annotations

import json
import logging
import secrets
import sys
import threading
import time
from datetime import datetime
from io import BytesIO
from pathlib import Path
from urllib.parse import quote

from fastapi import (
    FastAPI, HTTPException, Request,
    File, Form, UploadFile,
)
from fastapi.responses import (
    HTMLResponse, RedirectResponse, JSONResponse,
)
from pydantic import BaseModel

# ============================================================================
#  ПУТИ И ЛОГИРОВАНИЕ
# ============================================================================

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
    datefmt="%H:%M:%S",
)
log = logging.getLogger("serve")


# ============================================================================
#  ИЗВЛЕЧЕНИЕ ТЕКСТА ИЗ РАЗНЫХ ФОРМАТОВ
# ============================================================================

def extract_text_from_docx_bytes(data: bytes) -> str:
    """Читает .docx из байтов → текст."""
    from docx import Document
    doc = Document(BytesIO(data))
    parts = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            parts.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append("; ".join(cells))
    return "\n".join(parts)


def extract_text_from_pdf_bytes(data: bytes) -> str:
    """Читает .pdf из байтов → текст (нужен pypdf)."""
    try:
        from pypdf import PdfReader
    except ImportError:
        raise RuntimeError("Установите: python3 -m pip install pypdf")
    reader = PdfReader(BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def extract_text_from_txt_bytes(data: bytes) -> str:
    """Декодирует .txt с автоопределением кодировки."""
    for enc in ("utf-8", "cp1251", "koi8-r"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract_text(filename: str, data: bytes) -> str:
    """Диспетчер по расширению файла."""
    name = (filename or "").lower()
    if name.endswith(".docx"):
        return extract_text_from_docx_bytes(data)
    if name.endswith(".pdf"):
        return extract_text_from_pdf_bytes(data)
    if name.endswith(".txt"):
        return extract_text_from_txt_bytes(data)
    raise HTTPException(400, f"Формат не поддерживается: {filename}")


# ============================================================================
#  ПРИЛОЖЕНИЕ
# ============================================================================

app = FastAPI(title="ArtKnit Local")


# ============================================================================
#  ПОЛЬЗОВАТЕЛИ
#  ⚠️ Для демо — plain-text пароли. В проде: bcrypt/argon2 + БД.
# ============================================================================

USERS: dict[str, dict] = {
    # --- Врачи ---
    "doctor": {
        "password": "doctor123",
        "role": "doctor",
        "name": "Туаева А.С.",
        "doctor_name": "Туаева А.С.",
    },
    "doctor2": {
        "password": "doctor456",
        "role": "doctor",
        "name": "Кузнецова А.В.",
        "doctor_name": "Кузнецова А.В.",
    },
    # --- Пациенты ---
    "patient": {
        "password": "patient123",
        "role": "patient",
        "name": "Орлова Елена Петровна",
        "patient_card_id": "1030",
    },
    "patient2": {
        "password": "patient456",
        "role": "patient",
        "name": "Николаева Ольга Владимировна",
        "patient_card_id": "1032",
    },
    # --- Руководитель ---
    "manager": {
        "password": "manager123",
        "role": "manager",
        "name": "Главврач",
        "doctor_name": "",
    },
    # --- Администратор ---
    "admin": {
        "password": "admin123",
        "role": "admin",
        "name": "Администратор",
        "doctor_name": "",
    },
}

SESSIONS: dict[str, dict] = {}


# ============================================================================
#  АУТЕНТИФИКАЦИЯ — вспомогательные функции
# ============================================================================

def current_user(request: Request) -> dict | None:
    """Возвращает пользователя или None. Не бросает исключение."""
    token = request.cookies.get("artknit_token")
    if not token:
        auth = request.headers.get("authorization", "")
        if auth.startswith("Bearer "):
            token = auth[7:]
    if token and token in SESSIONS:
        return SESSIONS[token]
    return None


def require_role(request: Request, *roles: str) -> dict:
    """Бросает 401/403, если пользователь не подходит."""
    user = current_user(request)
    if not user:
        next_url = str(request.url.path)
        raise HTTPException(
            status_code=401,
            detail=f"Требуется вход. Перейдите на /login?next={quote(next_url)}",
        )
    if user["role"] not in roles:
        raise HTTPException(403, f"Доступ только для: {', '.join(roles)}")
    return user


def _default_redirect(role: str, user: dict | None = None) -> str:
    """Куда по умолчанию после логина."""
    if role == "doctor":
        return "/doctor/dashboard"
    if role == "patient":
        return "/patient/dashboard"
    if role == "manager":
        return "/manager"
    if role == "admin":
        return "/admin"
    return "/upload"


# ============================================================================
#  ИНЪЕКЦИЯ КНОПКИ «ВЫЙТИ»
#  ⚠️ Используем .replace(), а не .format() — иначе CSS-скобки { } ломают
#     форматирование (KeyError: ' padding-top').
# ============================================================================

_LOGOUT_BUTTON_HTML = """
<style>
  /* Отступ сверху, чтобы фиксированная плашка не перекрывала контент */
  body { padding-top: 64px !important; }
  @media (max-width: 520px) {
    body { padding-top: 100px !important; }
  }
</style>
<div id="artknit-logout" style="
    position:fixed; top:12px; right:16px; z-index:9999;
    display:flex; gap:8px; align-items:center;
    font-family:-apple-system,'Segoe UI',Roboto,sans-serif;
    font-size:12px;
">
    <span style="
        background:#fff; padding:6px 10px; border-radius:8px;
        color:#6B7280; box-shadow:0 2px 6px rgba(0,0,0,0.08);
        border:1px solid #E5E7EB; white-space:nowrap;
    ">
        <b style="color:#111827;">__NAME__</b> · __ROLE__
    </span>
    <a href="/api/logout" style="
        display:inline-flex; align-items:center; gap:5px;
        padding:6px 12px; background:#fff; color:#DC2626;
        border:1px solid #FCA5A5; border-radius:8px;
        text-decoration:none; font-weight:600; white-space:nowrap;
        box-shadow:0 2px 6px rgba(220,38,38,0.08);
        transition:all .15s;
    " onmouseover="this.style.background='#FEF2F2';"
       onmouseout="this.style.background='#fff';">
        <span>🚪</span> Выйти
    </a>
</div>
"""

_ROLE_LABELS = {
    "doctor":  "Врач",
    "patient": "Пациент",
    "manager": "Руководитель",
    "admin":   "Администратор",
}


def with_logout_button(html: str, user: dict | None) -> str:
    """Вставляет плавающую кнопку «Выйти» в HTML перед </body>."""
    if not user:
        return html

    name = user.get("name") or user.get("username") or "Пользователь"
    role = _ROLE_LABELS.get(user.get("role"), user.get("role", ""))

    # ⚠️ НЕ .format() — CSS-скобки ломают его. Используем .replace().
    button = (
        _LOGOUT_BUTTON_HTML
        .replace("__NAME__", str(name))
        .replace("__ROLE__", str(role))
    )

    if "</body>" in html:
        return html.replace("</body>", button + "</body>", 1)
    return html + button


# ============================================================================
#  МОДЕЛИ
# ============================================================================

class LoginIn(BaseModel):
    username: str
    password: str


# ============================================================================
#  ПУБЛИЧНЫЕ СТРАНИЦЫ
# ============================================================================

@app.get("/")
def root():
    """Главная → публичная загрузка."""
    return RedirectResponse("/upload")


@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request):
    """Страница входа. Поддерживает ?next=/doctor/auto_123."""
    template = (ROOT / "templates" / "login_static.html").read_text(encoding="utf-8")

    next_url = request.query_params.get("next", "")

    # Уже залогинен — сразу редирект
    user = current_user(request)
    if user:
        target = next_url or _default_redirect(user["role"], user)
        return RedirectResponse(target)

    if next_url:
        template = template.replace(
            "// __NEXT_URL_PLACEHOLDER__",
            f'window.NEXT_URL = {json.dumps(next_url)};'
        )
    return template


@app.get("/upload", response_class=HTMLResponse)
def upload_page():
    """Публичная страница загрузки. Авторизация НЕ требуется."""
    return (ROOT / "templates" / "upload.html.j2").read_text(encoding="utf-8")


# ============================================================================
#  API АУТЕНТИФИКАЦИИ
# ============================================================================

@app.post("/api/login")
def login(payload: LoginIn):
    """Логин по паре логин/пароль для всех ролей."""
    username = (payload.username or "").strip().lower()
    password = payload.password or ""

    entry = USERS.get(username)
    if not entry or entry["password"] != password:
        raise HTTPException(401, "Неверный логин или пароль")

    user = {
        "username": username,
        "role": entry["role"],
        "name": entry["name"],
        "doctor_name": entry.get("doctor_name", ""),
        "patient_card_id": entry.get("patient_card_id", ""),
    }

    token = secrets.token_urlsafe(32)
    SESSIONS[token] = user

    response = JSONResponse({
        "access_token": token,
        "role": user["role"],
        "name": user["name"],
    })
    response.set_cookie(
        key="artknit_token",
        value=token,
        httponly=False,   # для демо; в проде — True
        samesite="lax",
        max_age=8 * 3600,
        path="/",
    )
    log.info(f"Вход: {username} → {user['role']}")
    return response


@app.get("/api/me")
def me(request: Request):
    user = current_user(request)
    if not user:
        return {"authenticated": False}
    return {
        "authenticated": True,
        "role": user["role"],
        "name": user["name"],
        "username": user["username"],
    }


@app.get("/api/logout")
def logout():
    """Удаляет cookie и редиректит на /upload."""
    response = RedirectResponse("/upload", status_code=303)
    response.delete_cookie("artknit_token", path="/")
    log.info("Выход пользователя")
    return response


# ============================================================================
#  API ЗАГРУЗКИ — публичный
# ============================================================================

@app.post("/api/upload")
async def upload_protocol(
    file: UploadFile = File(...),
    patient_name: str = Form(...),
    patient_age: int = Form(...),
    study_type: str = Form(...),
    study_date: str = Form(""),
):
    """Публичный эндпоинт загрузки протокола."""
    # 1. Проверка типа файла
    allowed = (".docx", ".txt", ".pdf")
    if not (file.filename or "").lower().endswith(allowed):
        raise HTTPException(400, f"Поддерживаются только {allowed}")

    # 2. Извлечение текста
    content = await file.read()
    try:
        text = extract_text(file.filename, content)
    except HTTPException:
        raise
    except Exception as e:
        log.exception("Ошибка извлечения текста")
        raise HTTPException(400, f"Не удалось прочитать файл: {e}")

    if not text.strip():
        raise HTTPException(400, "Файл пуст или текст не извлёкся")

    log.info(f"Загрузка (публичная): {file.filename}, {len(text)} символов")

    # 3. Пайплайн
    from config_loader import ConfigLoader
    from extractor import Extractor
    from router import Router
    from generator import Generator

    cfg = ConfigLoader(ROOT / "config").load()
    ext = Extractor(cfg)
    rt = Router(cfg)
    gen = Generator(templates_dir=ROOT / "templates", output_dir=ROOT / "output")

    extraction = ext.extract(text)
    route = rt.build(extraction, study_type=study_type, study_date=study_date)

    # 4. Метаданные
    card_id = f"auto_{time.time_ns()}"
    patient = {
        "name": patient_name,
        "initials": "".join(p[0] for p in patient_name.split()[:2]).upper() or "—",
        "age": patient_age,
        "card_id": card_id,
        "sex_label": "—",
    }
    visit = {"date": study_date or "—", "time": "—", "doctor": "—"}
    doctor = {"name": "—", "initials": "—"}
    meta = {
        "sentences_total": extraction.sentences_total,
        "positive_count": len(extraction.positive()),
        "negative_count": len(extraction.negative()),
        "organ_codes": extraction.organ_codes or ["—"],
        "unknown_count": len(route.unknown_sentences),
    }

    # 5. Генерация HTML
    doctor_file = gen.render_doctor(route, patient, visit, meta, doctor)
    patient_file = gen.render_patient(route, patient, visit, meta)
    log.info(f"Созданы: {doctor_file.name}, {patient_file.name}")

    # ========================================================================
    #  6. Определяем, что показывать в модалке после обработки
    # ========================================================================
    specialists_lower = [s.lower() for s in route.specialist]
    is_oncologist = any("онколог" in s for s in specialists_lower)
    has_unknown = len(route.unknown_sentences) > 0

    # Требуется врач, если:
    #   • направили к онкологу И есть неопознанные находки
    #   • ИЛИ статус pending_doctor
    requires_doctor = (
        (is_oncologist and has_unknown)
        or route.status == "pending_doctor"
    )
    notification_sent = not requires_doctor

    log.info(
        f"Решение: oncologist={is_oncologist}, unknown={has_unknown}, "
        f"requires_doctor={requires_doctor}, notification_sent={notification_sent}"
    )

    # 7. Аудит
    audit_record = {
        "patient_id": card_id,
        "patient_name": patient_name,
        "patient_age": patient_age,
        "finding": route.finding_summary,
        "finding_id": route.primary_finding.id if route.primary_finding else None,
        "specialist": route.specialist,
        "urgency": route.urgency,
        "status": route.status,
        "status_label": route.status_label,
        "scenario": route.scenario,
        "multidisciplinary": route.multidisciplinary,
        "unknown_sentences": route.unknown_sentences,
        "uploaded_at": datetime.now().isoformat(),
        "source_file": file.filename,
        "organ_codes": meta["organ_codes"],
        "upload_mode": "public",
        "notification_sent": notification_sent,
        "requires_doctor_confirmation": requires_doctor,
    }
    audit_file = ROOT / "output" / "routes_audit.jsonl"
    audit_file.parent.mkdir(parents=True, exist_ok=True)
    with open(audit_file, "a", encoding="utf-8") as f:
        f.write(json.dumps(audit_record, ensure_ascii=False) + "\n")

    # 8. Автопересборка дашборда руководителя в фоне (не блокирует ответ)
    try:
        import subprocess
        subprocess.Popen(
            [sys.executable, str(ROOT / "scripts" / "build_manager.py")],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            cwd=str(ROOT),
        )
        log.info("Пересборка manager.html запущена в фоне")
    except Exception as e:
        log.warning(f"Не удалось пересобрать дашборд: {e}")

    # 9. Ответ — с флагами для модалки
    return {
        "ok": True,
        "route_id": card_id,
        "finding": route.finding_summary,
        "specialist": route.specialist,
        "urgency": route.urgency,
        "status": route.status,
        "status_label": route.status_label,
        "unknown_count": len(route.unknown_sentences),
        "doctor_html": f"/doctor/{card_id}",
        "patient_html": f"/patient/{card_id}",
        "doctor_file": doctor_file.name,
        "patient_file": patient_file.name,
        # флаги для модалки
        "is_oncologist": is_oncologist,
        "requires_doctor_confirmation": requires_doctor,
        "notification_sent": notification_sent,
        "notification_channel": "SMS" if notification_sent else None,
        "patient_name": patient_name,
    }


# ============================================================================
#  ПАЦИЕНТ — ЛИЧНЫЙ КАБИНЕТ
#  ⚠️ ОБЯЗАТЕЛЬНО до /patient/{card_id}
# ============================================================================

@app.get("/patient/dashboard", response_class=HTMLResponse)
def patient_dashboard(request: Request):
    """
    Личный кабинет пациента. Показывает ТОЛЬКО его маршрут.
    Привязка через patient_card_id в профиле пользователя.
    """
    user = require_role(request, "patient", "admin")

    card_id = user.get("patient_card_id", "")

    # Если прямого card_id нет — ищем по ФИО в аудите
    if not card_id:
        audit = ROOT / "output" / "routes_audit.jsonl"
        if audit.exists():
            for line in audit.read_text(encoding="utf-8").splitlines():
                try:
                    rec = json.loads(line)
                except Exception:
                    continue
                if rec.get("patient_name") == user["name"]:
                    card_id = rec.get("patient_id")
                    if card_id:
                        break

    if not card_id:
        html = f"""<html><head><meta charset="utf-8"><title>Кабинет пациента</title>
        <style>body{{font-family:-apple-system,sans-serif;padding:40px;
        max-width:600px;margin:0 auto;color:#111827;line-height:1.7;}}
        h1{{margin-bottom:16px;}} a{{color:#2563EB;}}</style></head><body>
        <h1>Маршрут пока не готов</h1>
        <p>Пациент: <b>{user['name']}</b></p>
        <p>Как только протокол будет обработан — он появится здесь.</p>
        </body></html>"""
        return HTMLResponse(with_logout_button(html, user))

    f = ROOT / "output" / f"patient_{card_id}.html"
    if not f.exists():
        html = f"""<html><head><meta charset="utf-8"><title>Кабинет пациента</title>
        <style>body{{font-family:-apple-system,sans-serif;padding:40px;
        max-width:600px;margin:0 auto;color:#111827;line-height:1.7;}}
        h1{{margin-bottom:16px;}} a{{color:#2563EB;}}</style></head><body>
        <h1>Маршрут пока не готов</h1>
        <p>Пациент: <b>{user['name']}</b></p>
        <p>Ожидаемый файл: <code>{f.name}</code></p>
        </body></html>"""
        return HTMLResponse(with_logout_button(html, user))

    return with_logout_button(f.read_text(encoding="utf-8"), user)


# ============================================================================
#  ПАЦИЕНТ — ПУБЛИЧНЫЙ MAGIC-LINK
#  ⚠️ ПОСЛЕ /patient/dashboard
# ============================================================================

@app.get("/patient/{card_id}", response_class=HTMLResponse)
def patient_screen(card_id: str):
    """
    Экран пациента по прямой ссылке.
    Доступен БЕЗ авторизации — пациент получает ссылку в SMS.
    """
    f = ROOT / "output" / f"patient_{card_id}.html"
    if not f.exists():
        raise HTTPException(404, f"Экран пациента не найден: {f.name}")
    return f.read_text(encoding="utf-8")


# ============================================================================
#  ВРАЧ — ЗАЩИЩЁННЫЕ СТРАНИЦЫ
#  ⚠️ Конкретные пути ДО параметрического /doctor/{card_id}
# ============================================================================

@app.get("/doctor/dashboard", response_class=HTMLResponse)
def doctor_dashboard(request: Request):
    """Список пациентов врача (последний сгенерированный экран)."""
    user = require_role(request, "doctor", "admin")
    files = sorted(
        (ROOT / "output").glob("doctor_*.html"),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    if not files:
        html = ("<h1>Нет данных</h1>"
                "<p>Загрузите протокол: <a href='/upload'>/upload</a></p>")
        return HTMLResponse(with_logout_button(html, user))
    return with_logout_button(files[0].read_text(encoding="utf-8"), user)


@app.get("/doctor/upload", response_class=HTMLResponse)
def doctor_upload_page(request: Request):
    """Врачебная загрузка — та же форма, но с авторизацией."""
    user = require_role(request, "doctor", "admin")
    log.info(f"Врачебная загрузка: {user['name']}")
    return (ROOT / "templates" / "upload.html.j2").read_text(encoding="utf-8")


@app.get("/doctor/{card_id}", response_class=HTMLResponse)
def doctor_screen(card_id: str, request: Request):
    """
    Экран врача для конкретного пациента. ТРЕБУЕТ авторизации.
    Если не залогинен — редирект на /login?next=/doctor/{card_id}.
    """
    user = current_user(request)
    if not user:
        next_url = f"/doctor/{card_id}"
        return RedirectResponse(f"/login?next={quote(next_url)}")
    if user["role"] not in ("doctor", "admin"):
        raise HTTPException(403, "Доступ только для врачей и администраторов")

    f = ROOT / "output" / f"doctor_{card_id}.html"
    if not f.exists():
        raise HTTPException(
            404,
            f"Экран врача не найден: {f.name}. Проверьте логи /api/upload.",
        )
    return with_logout_button(f.read_text(encoding="utf-8"), user)


# ============================================================================
#  МЕНЕДЖЕР / АДМИН
# ============================================================================

@app.get("/manager", response_class=HTMLResponse)
def manager_page(request: Request):
    user = require_role(request, "manager", "admin")
    f = ROOT / "output" / "manager.html"
    if not f.exists():
        html = ("<h1>Дашборд не построен</h1>"
                "<p>Запустите: <code>python3 scripts/build_manager.py</code></p>")
        return HTMLResponse(with_logout_button(html, user))
    return with_logout_button(f.read_text(encoding="utf-8"), user)


@app.get("/admin", response_class=HTMLResponse)
def admin_page(request: Request):
    user = require_role(request, "admin")
    html = """
        <html><head><meta charset="utf-8"><title>Админ</title>
        <style>
          body{font-family:-apple-system,sans-serif;padding:40px;
               max-width:700px;margin:0 auto;color:#111827;line-height:1.7;}
          h1{margin-bottom:16px;}
          a{color:#2563EB;text-decoration:none;}
          a:hover{text-decoration:underline;}
          ul{list-style:none;padding:0;}
          li{margin:10px 0;font-size:15px;}
        </style></head><body>
        <h1>Панель администратора</h1>
        <p>Полный доступ.</p>
        <ul>
            <li>🌐 <a href="/upload">Публичная загрузка</a></li>
            <li>👨‍⚕️ <a href="/doctor/dashboard">Кабинет врача</a></li>
            <li>🧑 <a href="/patient/dashboard">Кабинет пациента</a></li>
            <li>📊 <a href="/manager">Дашборд руководителя</a></li>
            <li>📚 <a href="/docs">API-документация</a></li>
        </ul>
        </body></html>
    """
    return HTMLResponse(with_logout_button(html, user))


# ============================================================================
#  ЗАПУСК
# ============================================================================

if __name__ == "__main__":
    import uvicorn
    import webbrowser

    def open_browser():
        time.sleep(1.5)
        try:
            webbrowser.open("http://127.0.0.1:8000/upload")
        except Exception:
            pass

    threading.Thread(target=open_browser, daemon=True).start()
    uvicorn.run(app, host="127.0.0.1", port=8000, log_level="info")
